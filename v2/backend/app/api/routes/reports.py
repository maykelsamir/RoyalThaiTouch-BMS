from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.db.session import get_db
from app.models.finance import Branch, DailyRevenue
from app.models.user import User
from app.schemas.reports import FinancialReportView, ReportBranchOption, ReportBranchSummary, ReportDailyRow
from app.services.effective_expenses import effective_amounts

router = APIRouter(prefix="/reports", tags=["reports"])


def _require(user: User, permission: str = "reports.view") -> None:
    if user.role.lower() != "admin" and permission not in user.permissions:
        raise HTTPException(status_code=403, detail="Reports permission required")


def _branch_ids(user: User, requested: list[int] | None, db: Session) -> list[int]:
    active = list(db.scalars(select(Branch.id).where(Branch.active.is_(True))))
    allowed = active if user.role.lower() == "admin" or not user.allowed_branch_ids else [x for x in active if x in user.allowed_branch_ids]
    if requested:
        if any(x not in allowed for x in requested):
            raise HTTPException(status_code=403, detail="Branch access denied")
        return list(dict.fromkeys(requested))
    return allowed


def _months(start: date, end: date):
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        yield year, month
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)


def _allocated_month(daily: Decimal, year: int, month: int, start: date, end: date) -> dict[date, Decimal]:
    cursor = max(start, date(year, month, 1))
    active_end = min(end, date(year, month, monthrange(year, month)[1]))
    values: dict[date, Decimal] = {}
    while cursor <= active_end and daily > 0:
        values[cursor] = daily
        cursor += timedelta(days=1)
    return values


def _percentage(value) -> Decimal:
    return Decimal(value if value is not None else 0)


def _share(revenue: Decimal, percentage: Decimal) -> Decimal:
    return (revenue * percentage / Decimal(100)).quantize(Decimal("1"))


def _report(start: date, end: date, ids: list[int], include_unapproved: bool, db: Session) -> FinancialReportView:
    if end < start:
        raise HTTPException(status_code=422, detail="End date must be after start date")
    if (end - start).days > 730:
        raise HTTPException(status_code=422, detail="Report range cannot exceed two years")

    branches = list(db.scalars(select(Branch).where(Branch.id.in_(ids)).order_by(Branch.name))) if ids else []
    query = select(DailyRevenue).where(DailyRevenue.branch_id.in_(ids), DailyRevenue.business_date.between(start, end))
    if not include_unapproved:
        query = query.where(DailyRevenue.status == "approved")
    revenues = list(db.scalars(query)) if ids else []
    revenue_map = {(x.branch_id, x.business_date): x for x in revenues}
    periods = list(_months(start, end))
    expense_map = effective_amounts(db, ids, periods)

    daily_rows, summaries = [], []
    gross_total = company_total = hotel_total = expense_total = Decimal(0)
    customer_total = 0

    for branch in branches:
        company_pct = _percentage(branch.company_revenue_percentage)
        hotel_pct = _percentage(branch.hotel_revenue_percentage)
        allocation: dict[date, Decimal] = {}
        for year, month in periods:
            allocation.update(_allocated_month(expense_map.get((branch.id, year, month), Decimal(0)), year, month, start, end))

        gross_sum = company_sum = hotel_sum = Decimal(0)
        customer_sum = approved = missing = 0
        cursor = start
        while cursor <= end:
            entry = revenue_map.get((branch.id, cursor))
            gross = Decimal(entry.amount or 0) if entry else Decimal(0)
            customers = int(entry.customer_count or 0) if entry else 0
            company_share = _share(gross, company_pct)
            hotel_share = _share(gross, hotel_pct)
            expense = allocation.get(cursor, Decimal(0))
            status = entry.status if entry else "missing"
            daily_rows.append(ReportDailyRow(
                business_date=cursor, branch_id=branch.id, branch_name=branch.name,
                revenue=gross, company_percentage=company_pct, hotel_percentage=hotel_pct,
                company_share=company_share, hotel_share=hotel_share, customer_count=customers,
                allocated_expense=expense, net_profit=company_share - expense, entry_status=status,
            ))
            gross_sum += gross
            company_sum += company_share
            hotel_sum += hotel_share
            customer_sum += customers
            approved += int(bool(entry and entry.status == "approved"))
            missing += int(entry is None)
            cursor += timedelta(days=1)

        expenses = sum(allocation.values(), Decimal(0))
        summaries.append(ReportBranchSummary(
            branch_id=branch.id, branch_name=branch.name, revenue=gross_sum,
            company_percentage=company_pct, hotel_percentage=hotel_pct,
            company_share=company_sum, hotel_share=hotel_sum, customer_count=customer_sum,
            expenses=expenses, net_profit=company_sum - expenses,
            approved_entries=approved, missing_days=missing,
        ))
        gross_total += gross_sum
        company_total += company_sum
        hotel_total += hotel_sum
        expense_total += expenses
        customer_total += customer_sum

    return FinancialReportView(
        date_from=start, date_to=end, generated_at=datetime.now(timezone.utc).isoformat(),
        company_revenue=gross_total, company_share=company_total, hotel_share=hotel_total,
        company_customer_count=customer_total, company_expenses=expense_total,
        company_net_profit=company_total - expense_total, branches=summaries, daily_rows=daily_rows,
    )


@router.get("/branches", response_model=list[ReportBranchOption])
def branches(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _require(current_user)
    ids = _branch_ids(current_user, None, db)
    rows = list(db.scalars(select(Branch).where(Branch.id.in_(ids)).order_by(Branch.name))) if ids else []
    return [ReportBranchOption(id=x.id, name=x.name) for x in rows]


@router.get("", response_model=FinancialReportView)
def view_report(date_from: date, date_to: date, branch_ids: list[int] | None = Query(default=None), include_unapproved: bool = False, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _require(current_user)
    return _report(date_from, date_to, _branch_ids(current_user, branch_ids, db), include_unapproved, db)


def _name(report: FinancialReportView, extension: str) -> str:
    return f"financial_report_{report.date_from}_{report.date_to}.{extension}"


def _money(value: Decimal) -> str:
    return f"{value:,.0f} IQD"


def _append_optional(headers: list, row: list, enabled: bool, header: str, value) -> None:
    if enabled:
        headers.append(header)
        row.append(value)


@router.get("/export/excel")
def excel(date_from: date, date_to: date, branch_ids: list[int] | None = Query(default=None), include_unapproved: bool = False, show_revenue: bool = True, show_expenses: bool = True, show_profit: bool = True, show_revenue_sharing: bool = True, show_branch_summary: bool = True, show_daily_details: bool = True, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _require(current_user, "reports.export_excel")
    report = _report(date_from, date_to, _branch_ids(current_user, branch_ids, db), include_unapproved, db)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    dark, red = "073C46", "D62828"
    summary.append(["Royal Thai Touch ERP - Financial Report"])
    summary.merge_cells("A1:J1")
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=dark)
    summary["A1"].alignment = Alignment(horizontal="center")
    summary.append(["Period", str(report.date_from), "to", str(report.date_to)])

    totals_headers, totals_values = [], []
    _append_optional(totals_headers, totals_values, show_revenue, "Gross Revenue", float(report.company_revenue))
    if show_revenue_sharing:
        totals_headers += ["Company Share", "Hotel Share"]
        totals_values += [float(report.company_share), float(report.hotel_share)]
    totals_headers.append("Customers")
    totals_values.append(report.company_customer_count)
    _append_optional(totals_headers, totals_values, show_expenses, "Expenses", float(report.company_expenses))
    _append_optional(totals_headers, totals_values, show_profit, "Company Net Profit", float(report.company_net_profit))
    summary.append(totals_headers)
    summary.append(totals_values)
    summary.append([])

    if show_branch_summary:
        headers = ["Branch"]
        if show_revenue: headers.append("Gross Revenue")
        if show_revenue_sharing: headers += ["Company %", "Hotel %", "Company Share", "Hotel Share"]
        headers.append("Customers")
        if show_expenses: headers.append("Expenses")
        if show_profit: headers.append("Company Net Profit")
        headers += ["Approved Days", "Missing Days"]
        summary.append(headers)
        for cell in summary[summary.max_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=dark)
        profit_col = headers.index("Company Net Profit") + 1 if show_profit else None
        for item in report.branches:
            row = [item.branch_name]
            if show_revenue: row.append(float(item.revenue))
            if show_revenue_sharing: row += [float(item.company_percentage), float(item.hotel_percentage), float(item.company_share), float(item.hotel_share)]
            row.append(item.customer_count)
            if show_expenses: row.append(float(item.expenses))
            if show_profit: row.append(float(item.net_profit))
            row += [item.approved_entries, item.missing_days]
            summary.append(row)
            if show_profit and item.net_profit < 0:
                cell = summary.cell(summary.max_row, profit_col)
                cell.fill = PatternFill("solid", fgColor=red)
                cell.font = Font(bold=True, color="FFFFFF")

    if show_revenue_sharing:
        sharing = workbook.create_sheet("Revenue Sharing")
        headers = ["Branch"]
        if show_revenue: headers.append("Gross Revenue")
        headers += ["Company %", "Hotel %", "Company Share", "Hotel Share"]
        if show_expenses: headers.append("Expenses")
        if show_profit: headers.append("Company Net Profit")
        sharing.append(headers)
        for cell in sharing[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=dark)
        for item in report.branches:
            row = [item.branch_name]
            if show_revenue: row.append(float(item.revenue))
            row += [float(item.company_percentage), float(item.hotel_percentage), float(item.company_share), float(item.hotel_share)]
            if show_expenses: row.append(float(item.expenses))
            if show_profit: row.append(float(item.net_profit))
            sharing.append(row)

    if show_daily_details:
        details = workbook.create_sheet("Daily Details")
        headers = ["Date", "Branch"]
        if show_revenue: headers.append("Gross Revenue")
        if show_revenue_sharing: headers += ["Company %", "Hotel %", "Company Share", "Hotel Share"]
        headers.append("Customers")
        if show_expenses: headers.append("Fixed Daily Expense")
        if show_profit: headers.append("Company Net Profit")
        headers.append("Status")
        details.append(headers)
        for cell in details[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=dark)
        profit_col = headers.index("Company Net Profit") + 1 if show_profit else None
        for item in report.daily_rows:
            row = [item.business_date, item.branch_name]
            if show_revenue: row.append(float(item.revenue))
            if show_revenue_sharing: row += [float(item.company_percentage), float(item.hotel_percentage), float(item.company_share), float(item.hotel_share)]
            row.append(item.customer_count)
            if show_expenses: row.append(float(item.allocated_expense))
            if show_profit: row.append(float(item.net_profit))
            row.append(item.entry_status)
            details.append(row)
            if show_profit and item.net_profit < 0:
                cell = details.cell(details.max_row, profit_col)
                cell.fill = PatternFill("solid", fgColor=red)
                cell.font = Font(bold=True, color="FFFFFF")
        details.freeze_panes = "A2"
        details.auto_filter.ref = details.dimensions

    for sheet in workbook.worksheets:
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(len(str(cell.value or "")) for cell in column) + 3, 38)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00' if "%" in str(sheet.cell(1, cell.column).value) else '#,##0'
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{_name(report, "xlsx")}"'})


@router.get("/export/pdf")
def pdf(date_from: date, date_to: date, branch_ids: list[int] | None = Query(default=None), include_unapproved: bool = False, show_revenue: bool = True, show_expenses: bool = True, show_profit: bool = True, show_revenue_sharing: bool = True, show_branch_summary: bool = True, show_daily_details: bool = True, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _require(current_user, "reports.export_pdf")
    report = _report(date_from, date_to, _branch_ids(current_user, branch_ids, db), include_unapproved, db)
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=8*mm, leftMargin=8*mm, topMargin=8*mm, bottomMargin=8*mm)
    styles = getSampleStyleSheet()
    dark, red = colors.HexColor("#073C46"), colors.HexColor("#D62828")
    story = [Paragraph("Royal Thai Touch ERP - Financial Report", styles["Title"]), Paragraph(f"Period: {report.date_from} to {report.date_to}", styles["Normal"]), Spacer(1, 5*mm)]

    headers, values = [], []
    _append_optional(headers, values, show_revenue, "Gross Revenue", _money(report.company_revenue))
    if show_revenue_sharing:
        headers += ["Company Share", "Hotel Share"]
        values += [_money(report.company_share), _money(report.hotel_share)]
    headers.append("Customers")
    values.append(f"{report.company_customer_count:,}")
    _append_optional(headers, values, show_expenses, "Expenses", _money(report.company_expenses))
    _append_optional(headers, values, show_profit, "Company Net Profit", _money(report.company_net_profit))
    totals = Table([headers, values], repeatRows=1)
    commands = [("BACKGROUND",(0,0),(-1,0),dark),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.5,colors.grey),("ALIGN",(1,0),(-1,-1),"RIGHT")]
    if show_profit and report.company_net_profit < 0:
        idx = headers.index("Company Net Profit")
        commands += [("BACKGROUND",(idx,1),(idx,1),red),("TEXTCOLOR",(idx,1),(idx,1),colors.white),("FONTNAME",(idx,1),(idx,1),"Helvetica-Bold")]
    totals.setStyle(TableStyle(commands))
    story += [totals, Spacer(1,5*mm)]

    if show_branch_summary:
        headers = ["Branch"]
        if show_revenue: headers.append("Gross Revenue")
        if show_revenue_sharing: headers += ["Co %", "Hotel %", "Company Share", "Hotel Share"]
        headers.append("Customers")
        if show_expenses: headers.append("Expenses")
        if show_profit: headers.append("Company Net")
        headers += ["Approved", "Missing"]
        rows = [headers]
        for item in report.branches:
            row = [item.branch_name]
            if show_revenue: row.append(_money(item.revenue))
            if show_revenue_sharing: row += [f"{item.company_percentage}%", f"{item.hotel_percentage}%", _money(item.company_share), _money(item.hotel_share)]
            row.append(f"{item.customer_count:,}")
            if show_expenses: row.append(_money(item.expenses))
            if show_profit: row.append(_money(item.net_profit))
            row += [item.approved_entries, item.missing_days]
            rows.append(row)
        table = Table(rows, repeatRows=1)
        commands = [("BACKGROUND",(0,0),(-1,0),dark),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.3,colors.grey),("FONTSIZE",(0,0),(-1,-1),6.8),("ALIGN",(1,1),(-1,-1),"RIGHT")]
        if show_profit:
            profit_col = headers.index("Company Net")
            for i, item in enumerate(report.branches, 1):
                if item.net_profit < 0:
                    commands += [("BACKGROUND",(profit_col,i),(profit_col,i),red),("TEXTCOLOR",(profit_col,i),(profit_col,i),colors.white)]
        table.setStyle(TableStyle(commands))
        story += [Paragraph("Branch Summary", styles["Heading2"]), table, Spacer(1,5*mm)]

    if show_daily_details:
        headers = ["Date", "Branch"]
        if show_revenue: headers.append("Gross")
        if show_revenue_sharing: headers += ["Co %", "Hotel %", "Company Share", "Hotel Share"]
        headers.append("Customers")
        if show_expenses: headers.append("Expense")
        if show_profit: headers.append("Company Net")
        headers.append("Status")
        rows = [headers]
        for item in report.daily_rows:
            row = [str(item.business_date), item.branch_name]
            if show_revenue: row.append(_money(item.revenue))
            if show_revenue_sharing: row += [f"{item.company_percentage}%", f"{item.hotel_percentage}%", _money(item.company_share), _money(item.hotel_share)]
            row.append(f"{item.customer_count:,}")
            if show_expenses: row.append(_money(item.allocated_expense))
            if show_profit: row.append(_money(item.net_profit))
            row.append(item.entry_status)
            rows.append(row)
        table = Table(rows, repeatRows=1)
        commands = [("BACKGROUND",(0,0),(-1,0),dark),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),6),("ALIGN",(2,1),(-2,-1),"RIGHT")]
        if show_profit:
            profit_col = headers.index("Company Net")
            for i, item in enumerate(report.daily_rows, 1):
                if item.net_profit < 0:
                    commands += [("BACKGROUND",(profit_col,i),(profit_col,i),red),("TEXTCOLOR",(profit_col,i),(profit_col,i),colors.white)]
        table.setStyle(TableStyle(commands))
        story += [Paragraph("Daily Details", styles["Heading2"]), table]

    document.build(story)
    output.seek(0)
    return StreamingResponse(output, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{_name(report, "pdf")}"'})
